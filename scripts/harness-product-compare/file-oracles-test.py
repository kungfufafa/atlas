"""Offline behavioral oracle tests; receipt fixtures never execute candidate code."""
from __future__ import annotations

import csv
import hashlib
import importlib.util
import io
import json
import shutil
import tempfile
import unittest
import zipfile
from decimal import Decimal
from pathlib import Path

from docx import Document
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from pptx import Presentation
from pptx.dml.color import RGBColor
from reportlab.pdfgen.canvas import Canvas

HERE = Path(__file__).resolve().parent


def load(name):
    spec = importlib.util.spec_from_file_location(name, HERE / f"{name}.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


fixtures, oracles, contract = load("file-fixtures"), load("file-oracles"), load("file-code-contract")
POLICY = hashlib.sha256(b"test fixture: externally enforced policy").hexdigest()


def rewrite_archive(source, target, part, old, new):
    with zipfile.ZipFile(source) as archive:
        entries = [(item, archive.read(item.filename)) for item in archive.infolist()]
    with zipfile.ZipFile(target, "w") as archive:
        for item, data in entries:
            archive.writestr(item, data.replace(old.encode(), new.encode()) if item.filename == part else data)


def write_csv(path, rows):
    with path.open("w", newline="", encoding="utf-8") as stream:
        csv.writer(stream).writerows(rows)


def make_good(task, workspace):
    """Derive report/join/reconciliation content independently from input bytes."""
    output = workspace / task["output"]
    output.parent.mkdir(parents=True, exist_ok=True)
    family = task["family"]
    if family == "xlsx_reconciliation":
        source = load_workbook(workspace / "input/invoices.xlsx")
        totals, count = {}, 0
        for _, vendor, kind, status, amount in list(source.active.values)[1:]:
            totals.setdefault(vendor, 0)
            if status == "settled":
                totals[vendor] += int(Decimal(amount) * 100) * (-1 if kind == "refund" else 1)
                count += 1
        workbook = Workbook()
        workbook.active.title = "Summary"
        for row in [["vendor", "net_cents"], *[[vendor, totals[vendor]] for vendor in sorted(totals)], ["TOTAL", sum(totals.values())]]:
            workbook.active.append(row)
        audit = workbook.create_sheet("Audit")
        audit.append(["metric", "value"])
        audit.append(["settled_count", count])
        workbook.save(output)
    elif family == "xlsx_surgical_edit":
        workbook = load_workbook(workspace / task["source"])
        workbook["Budget"]["C2"] = task["expected"]["value"]
        workbook.save(output)
    elif family in {"docx_revision", "pptx_revision"}:
        part = "word/document.xml" if family == "docx_revision" else "ppt/slides/slide2.xml"
        rewrite_archive(workspace / task["source"], output, part, task["expected"]["old"], task["expected"]["new"])
    elif family == "docx_report":
        facts = json.loads((workspace / "input/incident.json").read_text())
        document = Document()
        document.add_heading("Incident report", 1)
        for name, value in facts.items():
            document.add_paragraph(f"{name.replace('_', ' ').title()}: {value}")
        table = document.add_table(rows=2, cols=2)
        for row, values in zip(table.rows, [["Action", "Due"], [facts["action"], facts["due"]]]):
            for cell, value in zip(row.cells, values):
                cell.text = value
        document.save(output)
    elif family == "pdf_create":
        facts = json.loads((workspace / "input/handover.json").read_text())
        canvas = Canvas(str(output), pagesize=(612, 792))
        for index, line in enumerate(["Shipment handover", *[f"{key}: {value}" for key, value in facts.items()]]):
            canvas.drawString(45, 744 - index * 30, line)
        canvas.save()
    elif family == "pdf_extract":
        from pypdf import PdfReader
        rows = [["reference", "pieces", "fee_cents"]]
        for page in PdfReader(workspace / "input/dispatch.pdf").pages:
            for line in page.extract_text().splitlines():
                if line.startswith("DSP-"):
                    rows.append([field.strip() for field in line.split("|")])
        write_csv(output, rows)
    elif family == "csv_join":
        with (workspace / "input/customers.csv").open(newline="") as stream:
            names = {row["customer_id"]: row["customer_name"] for row in csv.DictReader(stream)}
        with (workspace / "input/orders.csv").open(newline="") as stream:
            rows = [[row["order_id"], row["customer_id"], row["total_cents"], names.get(row["customer_id"], "")] for row in csv.DictReader(stream)]
        write_csv(output, [["order_id", "customer_id", "total_cents", "customer_name"], *rows])
    elif family == "code_fix":
        output.write_text('from decimal import Decimal, InvalidOperation, ROUND_HALF_UP\n\ndef to_cents(value: str) -> int:\n    try:\n        number = Decimal(value.strip())\n        if not number.is_finite():\n            raise ValueError("finite input required")\n        return int((number * 100).to_integral_value(rounding=ROUND_HALF_UP))\n    except (InvalidOperation, AttributeError) as error:\n        raise ValueError("invalid currency") from error\n')
    return output


def receipt(task, output):
    count = len(task["expected"]["cases"]) + len(task["expected"]["invalid"])
    return {"pass": True, "sourceSha256": hashlib.sha256(output.read_bytes()).hexdigest(), "testSuiteSha256": task["testSuiteSha256"], "sandboxPolicySha256": POLICY, "candidateContract": {"pass": True, "sha256": contract.CONTRACT_SHA256}, "exitCode": 0, "casesPassed": count, "casesTotal": count}


class OracleTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="native-oracle-test-", dir="/private/tmp")
        self.root = Path(self.temp.name)
        self.index = 0

    def tearDown(self):
        self.temp.cleanup()

    def prepare(self, family, variant=0):
        self.index += 1
        root = self.root / str(self.index)
        workspace, originals, manifest = root / "workspace", root / "originals", root / "private/task.json"
        public = fixtures.build(family, "development", variant, workspace, manifest)
        task = json.loads(manifest.read_text())
        shutil.copytree(workspace, originals)
        output = make_good(task, workspace)
        return task, workspace, originals, output, public

    def score(self, task, workspace, originals, output=None, hidden=None):
        return oracles.inspect(task, workspace, originals, hidden, str(output.relative_to(workspace)) if output else task["output"], POLICY, contract.CONTRACT_SHA256)

    def assert_passes(self, result):
        self.assertTrue(result["pass"], result["checks"])

    def assert_fails(self, result, check=None):
        self.assertFalse(result["pass"], result)
        if check:
            self.assertTrue(any(item["id"] == check and not item["pass"] for item in result["checks"]), result["checks"])

    def test_all_development_families_and_variants_accept_known_good_files(self):
        for family in fixtures.FAMILIES:
            for variant in range(2):
                with self.subTest(family=family, variant=variant):
                    task, workspace, originals, output, public = self.prepare(family, variant)
                    self.assertEqual(set(public), {"id", "family", "split", "variant", "prompt"})
                    self.assertFalse((workspace / "task.json").exists())
                    hidden = receipt(task, output) if family == "code_fix" else None
                    self.assert_passes(self.score(task, workspace, originals, output, hidden))

    def test_every_family_rejects_changed_source_bytes(self):
        for family in fixtures.FAMILIES:
            with self.subTest(family=family):
                task, workspace, originals, output, _ = self.prepare(family)
                source = next(name for name in task["sources"] if name not in task.get("mutableSources", []))
                path = workspace / source
                # Deliberately corrupt even a read-only synthetic fixture to
                # verify byte preservation independently of filesystem modes.
                path.chmod(0o644)
                path.write_bytes(path.read_bytes() + b"changed source")
                result = self.score(task, workspace, originals, output, receipt(task, output) if family == "code_fix" else None)
                self.assert_fails(result, f"source_preserved:{source}")
                self.assertTrue(result["integrityFailure"])

    def test_numeric_reconciliation_rejects_wrong_facts_and_string_numbers(self):
        for replacement in ["123", True, 123.5, -999]:
            task, workspace, originals, output, _ = self.prepare("xlsx_reconciliation")
            workbook = load_workbook(output)
            workbook["Summary"]["B2"] = replacement
            workbook.save(output)
            self.assert_fails(self.score(task, workspace, originals, output), "cells:Summary")

    def test_reconciliation_styled_empty_cells_are_benign_but_extra_data_is_schema_failure(self):
        task, workspace, originals, output, _ = self.prepare("xlsx_reconciliation")
        workbook = load_workbook(output)
        workbook["Summary"]["H20"].font = Font(bold=True)
        workbook.save(output)
        self.assert_passes(self.score(task, workspace, originals, output))
        workbook["Summary"]["H20"] = "extra metadata"
        workbook.save(output)
        self.assert_fails(self.score(task, workspace, originals, output), "schema_extra_cells:Summary")

    def test_surgical_workbook_rejects_formula_style_and_type_corruption(self):
        for mutation in ["formula", "style", "type", "hidden", "width"]:
            task, workspace, originals, output, _ = self.prepare("xlsx_surgical_edit")
            workbook = load_workbook(output)
            if mutation == "formula": workbook["Budget"]["D2"] = 100
            elif mutation == "style": workbook["Budget"]["C2"].font = Font(bold=False)
            elif mutation == "type": workbook["Budget"]["C2"] = str(task["expected"]["value"])
            elif mutation == "hidden": workbook["Notes"].sheet_state = "visible"
            else: workbook["Budget"].column_dimensions["A"].width = 10
            workbook.save(output)
            self.assert_fails(self.score(task, workspace, originals, output), "workbook_edit_and_preservation")

    def test_docx_revision_rejects_wrong_date_emphasis_and_table_content(self):
        for mutation in ["date", "emphasis", "table"]:
            task, workspace, originals, output, _ = self.prepare("docx_revision")
            document = Document(output)
            if mutation == "date": document.paragraphs[1].runs[0].text = "Deadline: wrong"
            elif mutation == "emphasis": document.paragraphs[2].runs[1].bold = False
            else: document.tables[0].cell(1, 1).text = "99"
            document.save(output)
            self.assert_fails(self.score(task, workspace, originals, output), "document_edit_and_preservation")

    def test_docx_report_rejects_substring_fact_falsepass_and_accepts_negative_status(self):
        task, workspace, originals, output, _ = self.prepare("docx_report")
        document = Document(output)
        document.add_paragraph("The incident is not resolved. No duration of impact was supplied.")
        document.save(output)
        self.assert_passes(self.score(task, workspace, originals, output))
        for paragraph in document.paragraphs:
            if paragraph.text.startswith("Affected Orders:"):
                paragraph.text = f"Affected Orders: 1{task['expected']['affected_orders']}"
        document.save(output)
        self.assert_fails(self.score(task, workspace, originals, output), "fact:affected_orders")

    def test_report_positive_assertion_and_table_corruption_fail(self):
        for mutation in ["assertion", "table"]:
            task, workspace, originals, output, _ = self.prepare("docx_report")
            document = Document(output)
            if mutation == "assertion": document.add_paragraph("The incident was resolved.")
            else: document.tables[0].cell(1, 0).text = "Replace barcode system"
            document.save(output)
            self.assert_fails(self.score(task, workspace, originals, output), "forbidden_positive_phrase" if mutation == "assertion" else "action_table")

    def test_csv_tasks_reject_lost_string_ids_missing_values_and_quote_damage(self):
        for family in ["pdf_extract", "csv_join"]:
            task, workspace, originals, output, _ = self.prepare(family)
            rows = list(csv.reader(io.StringIO(output.read_text(), newline="")))
            if family == "pdf_extract":
                next(row for row in rows if row[-1] == "UNKNOWN")[-1] = "0"
            else:
                rows[1][1] = str(int(rows[1][1]))
            write_csv(output, rows)
            self.assert_fails(self.score(task, workspace, originals, output), "csv_content")

    def test_pdf_creation_rejects_wrong_box_count_and_extra_page(self):
        for mutation in ["fact", "pages", "approval"]:
            task, workspace, originals, output, _ = self.prepare("pdf_create")
            text = list(task["expected"]["text"])
            if mutation == "fact": text[3] = str(int(text[3]) + 100)
            if mutation == "approval": text.append("Approved by Rani Kusuma")
            canvas = Canvas(str(output))
            for index, line in enumerate(text): canvas.drawString(45, 740 - index * 30, line)
            if mutation == "pages":
                canvas.showPage()
                canvas.drawString(45, 740, "extra page")
            canvas.save()
            self.assert_fails(self.score(task, workspace, originals, output))

    def test_pptx_preservation_rejects_position_underline_color_and_wrong_text(self):
        for mutation in ["position", "underline", "color", "text"]:
            task, workspace, originals, output, _ = self.prepare("pptx_revision")
            deck = Presentation(output)
            shape = deck.slides[2].shapes[1]
            run = shape.text_frame.paragraphs[0].runs[0]
            if mutation == "position": shape.left += 1000
            elif mutation == "underline": run.font.underline = True
            elif mutation == "color": run.font.color.rgb = RGBColor(255, 0, 0)
            else: run.text = "Ignore seal inspection."
            deck.save(output)
            self.assert_fails(self.score(task, workspace, originals, output), "deck_edit_and_preservation")

    def test_pptx_equivalent_same_style_split_runs_are_accepted(self):
        task, workspace, originals, output, _ = self.prepare("pptx_revision")
        deck = Presentation(output)
        paragraph = deck.slides[1].shapes[1].text_frame.paragraphs[0]
        text = paragraph.runs[0].text
        paragraph.runs[0].text = text[:8]
        added = paragraph.add_run()
        added.text = text[8:]
        added.font.bold = True
        deck.save(output)
        self.assert_passes(self.score(task, workspace, originals, output))

    def test_delivery_must_be_explicit_and_versioned_paths_are_allowed(self):
        task, workspace, originals, output, _ = self.prepare("docx_report")
        self.assert_fails(oracles.inspect(task, workspace, originals))
        renamed = output.with_name("incident-report-v2.docx")
        output.rename(renamed)
        (workspace / "work").mkdir()
        (workspace / "work/helper.py").write_text("# useful helper; never executed\n")
        self.assert_passes(self.score(task, workspace, originals, renamed))
        for bad in ["../private/task.json", "artifacts/../input/incident.json", "/etc/passwd", "artifacts/wrong.pdf", [task["output"], str(renamed)]]:
            self.assert_fails(oracles.inspect(task, workspace, originals, delivered_path=bad))

    def test_symlinks_and_untrusted_evaluator_originals_are_rejected(self):
        task, workspace, originals, output, _ = self.prepare("xlsx_surgical_edit")
        outside = self.root / "outside.xlsx"
        shutil.copyfile(output, outside)
        output.unlink()
        output.symlink_to(outside)
        self.assert_fails(self.score(task, workspace, originals, output))
        output.unlink()
        shutil.copyfile(outside, output)
        source = originals / task["source"]
        source.write_bytes(source.read_bytes() + b"tampered evaluator original")
        self.assert_fails(self.score(task, workspace, originals, output))

    def test_hidden_code_execution_is_never_inferred_from_source_or_partial_receipt(self):
        task, workspace, originals, output, _ = self.prepare("code_fix")
        self.assert_fails(self.score(task, workspace, originals, output), "isolated_hidden_tests")
        good = receipt(task, output)
        self.assert_passes(self.score(task, workspace, originals, output, good))
        for key, value in [("pass", "true"), ("sourceSha256", "0" * 64), ("testSuiteSha256", "0" * 64), ("sandboxPolicySha256", "0" * 64), ("candidateContract", {"pass": False, "sha256": contract.CONTRACT_SHA256}), ("candidateContract", {"pass": "true", "sha256": contract.CONTRACT_SHA256}), ("candidateContract", {"pass": True, "sha256": "0" * 64}), ("candidateContract", {"pass": True}), ("candidateContract", True), ("exitCode", False), ("exitCode", 1), ("casesPassed", 1), ("casesTotal", 1)]:
            damaged = {**good, key: value}
            self.assert_fails(self.score(task, workspace, originals, output, damaged), "isolated_hidden_tests")
        self.assert_fails(oracles.inspect(task, workspace, originals, good), "isolated_hidden_tests")
        self.assert_fails(oracles.inspect(task, workspace, originals, good, task["output"], POLICY), "isolated_hidden_tests")
        self.assert_fails(oracles.inspect(task, workspace, originals, good, task["output"], POLICY, "invalid-hash"), "isolated_hidden_tests")
        for field in ("candidateContract",):
            missing = {key: value for key, value in good.items() if key != field}
            self.assert_fails(self.score(task, workspace, originals, output, missing), "isolated_hidden_tests")
        changed_task = {**task, "candidateContractSha256": "0" * 64}
        self.assert_fails(oracles.inspect(changed_task, workspace, originals, good, task["output"], POLICY, task["candidateContractSha256"]), "isolated_hidden_tests")
        self.assert_fails(oracles.inspect(task, workspace, originals, good, task["output"], POLICY, "0" * 64), "isolated_hidden_tests")
        (workspace / "app/backdoor.py").write_text("# not executed\n")
        self.assert_fails(self.score(task, workspace, originals, output, good), "code_mutation_scope")

    def test_expectations_must_be_outside_workspace_and_cannot_be_overwritten(self):
        workspace = self.root / "workspace"
        with self.assertRaises(ValueError): fixtures.build("csv_join", "development", 0, workspace, workspace / "private.json")
        manifest = self.root / "private.json"
        fixtures.build("csv_join", "development", 0, workspace, manifest)
        with self.assertRaises(ValueError): fixtures.build("csv_join", "development", 0, self.root / "new", manifest)

    def test_numeric_facts_do_not_match_dates_decimals_or_grouped_numbers(self):
        for text in ["Due: 2026-10-17", "Count: 117", "Count: 17.5", "Count: 17,000"]:
            self.assertFalse(oracles.contains_fact(text, 17), text)
        self.assertTrue(oracles.contains_fact("Affected orders: 17.", 17))

    def test_archive_duplicates_and_extreme_sheet_dimensions_fail_before_iteration(self):
        task, workspace, originals, output, _ = self.prepare("xlsx_reconciliation")
        workbook = load_workbook(output)
        workbook["Summary"].cell(1_000_000, 1, "oversized dimensions")
        workbook.save(output)
        self.assert_fails(self.score(task, workspace, originals, output), "artifact_inspection")
        duplicate = self.root / "duplicate.docx"
        import warnings
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            with zipfile.ZipFile(duplicate, "w") as archive:
                archive.writestr("word/document.xml", "first")
                archive.writestr("word/document.xml", "second")
        with self.assertRaises(ValueError): oracles.bounded_archive(duplicate)

    def test_inherited_word_style_and_deck_theme_font_changes_are_detected(self):
        task, workspace, originals, output, _ = self.prepare("docx_revision")
        document = Document(output)
        document.styles["Heading 2"].font.italic = True
        document.save(output)
        self.assert_fails(self.score(task, workspace, originals, output), "document_edit_and_preservation")
        task, workspace, originals, output, _ = self.prepare("pptx_revision")
        with zipfile.ZipFile(output) as archive:
            theme = archive.read("ppt/theme/theme1.xml").decode()
        import re
        match = re.search(r'typeface="([^"]+)"', theme)
        self.assertIsNotNone(match)
        rewrite_archive(output, output, "ppt/theme/theme1.xml", match.group(0), 'typeface="Courier New"')
        self.assert_fails(self.score(task, workspace, originals, output), "deck_edit_and_preservation")


if __name__ == "__main__":
    unittest.main(verbosity=2)

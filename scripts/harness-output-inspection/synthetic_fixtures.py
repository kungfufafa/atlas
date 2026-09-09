"""Build small, public synthetic inspection fixtures without running a renderer.

Bytes are deterministic for a fixed output path and dependency versions. The
external-link DOCX necessarily includes its owned canary's absolute file URI;
that one member changes when the fresh fixture root changes. No historical
fixture, user document, credential or study content is read.
"""

from datetime import datetime, timezone
import importlib.metadata
import io
from pathlib import Path
import zipfile

from docx import Document
from docx.shared import Inches
from lxml import etree as ET
from PIL import Image, ImageDraw
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas


FIXED_TIME = datetime(2000, 1, 1, tzinfo=timezone.utc)
FIXED_ZIP_TIME = (1980, 1, 1, 0, 0, 0)
RELATIONSHIPS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
DRAWING = "http://schemas.openxmlformats.org/drawingml/2006/main"


def dependency_versions():
    return {
        name: importlib.metadata.version(name)
        for name in ("python-docx", "Pillow", "reportlab", "lxml")
    }


def write_package(path, members):
    with Path(path).open("xb") as stream:
        with zipfile.ZipFile(stream, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as package:
            for name, data in sorted(members.items()):
                info = zipfile.ZipInfo(name, date_time=FIXED_ZIP_TIME)
                info.create_system = 3
                info.external_attr = 0o644 << 16
                info.compress_type = zipfile.ZIP_DEFLATED
                package.writestr(info, data, compresslevel=9)


def save_document(document, path):
    properties = document.core_properties
    properties.created = FIXED_TIME
    properties.modified = FIXED_TIME
    properties.author = "Synthetic inspection fixture"
    properties.last_modified_by = "Synthetic inspection fixture"
    properties.revision = 1
    properties.title = "Synthetic inspection fixture"
    raw = io.BytesIO()
    document.save(raw)
    with zipfile.ZipFile(io.BytesIO(raw.getvalue())) as package:
        members = {name: package.read(name) for name in package.namelist()}
    write_package(path, members)


def build_fixtures(root):
    root = Path(root)
    if not root.is_absolute():
        raise ValueError("Fixture root must be absolute")
    root.mkdir()  # Exclusive and retained, including any later generation failure.
    captured = root / "captured-inputs"
    captured.mkdir()
    outside = root / "outside-private"
    outside.mkdir()
    paths = {
        "embedded_docx": captured / "embedded.docx",
        "external_docx": captured / "external-link.docx",
        "complete_docx": captured / "complete.docx",
        "complete_pdf": captured / "complete.pdf",
        "malformed_docx": captured / "not-a-docx.docx",
        "xml_error_docx": captured / "xml-error.docx",
        "many_pages_pdf": captured / "too-many-pages.pdf",
        "public_image": captured / "public.png",
        "visible_image": captured / "visible.png",
        "red_image": outside / "private-image.png",
    }

    blue = Image.new("RGB", (1200, 200), "#125be3")
    ImageDraw.Draw(blue).text((30, 75), "PUBLIC EMBEDDED IMAGE", fill="white", font_size=45)
    blue.save(paths["public_image"], format="PNG")
    red = Image.new("RGB", (1200, 200), "#fc1842")
    ImageDraw.Draw(red).text((30, 75), "SYNTHETIC OUTSIDE IMAGE", fill="white", font_size=35)
    red.save(paths["red_image"], format="PNG")

    embedded = Document()
    embedded.add_paragraph("Public document, resource isolation fixture.")
    embedded.add_picture(str(paths["public_image"]), width=Inches(6))
    save_document(embedded, paths["embedded_docx"])
    with zipfile.ZipFile(paths["embedded_docx"]) as package:
        members = {name: package.read(name) for name in package.namelist() if not name.startswith("word/media/")}
    tree = ET.fromstring(members["word/document.xml"])
    for blip in tree.findall(".//{" + DRAWING + "}blip"):
        identity = blip.attrib.pop("{" + RELATIONSHIPS + "}embed")
        blip.attrib["{" + RELATIONSHIPS + "}link"] = identity
    members["word/document.xml"] = ET.tostring(tree, xml_declaration=True, encoding="UTF-8", standalone=True)
    tree = ET.fromstring(members["word/_rels/document.xml.rels"])
    for relationship in tree:
        if relationship.attrib.get("Type", "").endswith("/image"):
            relationship.attrib.update({"Target": paths["red_image"].as_uri(), "TargetMode": "External"})
    members["word/_rels/document.xml.rels"] = ET.tostring(tree, xml_declaration=True, encoding="UTF-8", standalone=True)
    write_package(paths["external_docx"], members)

    visible = Image.new("RGB", (640, 150), "white")
    ImageDraw.Draw(visible).text((20, 45), "Synthetic image-only visible text", fill="black")
    visible.save(paths["visible_image"], format="PNG")
    complete = Document()
    complete.add_heading("Synthetic inspection example", 0)
    complete.add_paragraph("Ordinary paragraph with punctuation.")
    complete.add_paragraph("Numbered item", style="List Number")
    table = complete.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Header one"
    table.cell(0, 1).text = "Header two"
    table.cell(1, 0).text = "Outer cell"
    nested = table.cell(1, 1).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "Nested cell content"
    complete.add_picture(str(paths["visible_image"]), width=Inches(5))
    complete.sections[0].header.paragraphs[0].text = "Synthetic running header"
    complete.sections[0].footer.paragraphs[0].text = "Synthetic running footer"
    hidden = complete.add_paragraph().add_run("Hidden synthetic text")
    hidden.font.hidden = True
    complete.add_page_break()
    complete.add_paragraph("Second page body.")
    save_document(complete, paths["complete_docx"])
    with zipfile.ZipFile(paths["complete_docx"]) as package:
        broken = {name: package.read(name) for name in package.namelist()}
    broken["word/document.xml"] = b"<invalid"
    write_package(paths["xml_error_docx"], broken)

    pdf = canvas.Canvas(str(paths["complete_pdf"]), pagesize=A4, invariant=1, pageCompression=1)
    pdf.drawString(50, 760, "Synthetic selectable PDF text")
    # Stored annotation only; fixture generation performs no HTTP request.
    pdf.linkURL("https://example.invalid/inspection", (50, 735, 300, 750))
    pdf.drawImage(ImageReader(visible), 50, 550, width=500, height=120)
    pdf.showPage()
    pdf.drawImage(ImageReader(visible), 50, 600, width=500, height=120)
    pdf.showPage()
    pdf.save()
    paths["malformed_docx"].write_bytes(b"plain bytes, no ZIP")
    many = canvas.Canvas(str(paths["many_pages_pdf"]), pagesize=A4, invariant=1, pageCompression=1)
    for page in range(33):
        many.drawString(50, 750, "Synthetic page " + str(page + 1))
        many.showPage()
    many.save()
    return paths

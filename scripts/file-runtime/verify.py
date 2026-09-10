"""Exercise Atlas's installed analysis libraries against real roundtrip files.

Run in a disposable workspace; writes fixtures under artifacts/runtime-check/.
"""
import csv
import importlib.metadata
import json
from pathlib import Path

import bidi
import fontTools
import uharfbuzz
import pandas as pd
from docx import Document
from openpyxl import Workbook, load_workbook
from PIL import Image, ImageChops, ImageStat
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen import canvas


def verify():
    output = Path("artifacts/runtime-check")
    output.mkdir(parents=True, exist_ok=True)

    source_path = output / "source.jpg"
    with Image.new("RGB", (32, 24), (20, 80, 120)) as source:
        source.paste((180, 40, 30), (0, 0, 16, 24))
        source.save(source_path, format="JPEG", quality=95)
    source_bytes = source_path.read_bytes()
    with Image.open(source_path) as source:
        source.load()
        edited = source.crop((0, 2, 24, 22)).resize((48, 40))
        with source.resize(edited.size) as resized_source:
            assert ImageChops.difference(edited, resized_source).getbbox() is not None
    with edited:
        for extension, image_format in (("jpg", "JPEG"), ("png", "PNG")):
            destination = output / f"edited.{extension}"
            edited.save(destination, format=image_format)
            with Image.open(destination) as reopened:
                reopened.load()
                assert reopened.format == image_format
                assert reopened.size == (48, 40)
                assert reopened.mode == "RGB"
                if image_format == "PNG":
                    assert reopened.tobytes() == edited.tobytes()
                else:
                    difference = ImageStat.Stat(ImageChops.difference(edited, reopened))
                    assert max(difference.mean) < 10
    assert source_path.read_bytes() == source_bytes

    doc = Document()
    doc.add_heading("Atlas daily report", 0)
    doc.add_paragraph("Jakarta: draft")
    doc.save(output / "source.docx")
    edited = Document(output / "source.docx")
    edited.paragraphs[1].text = "Jakarta: verified"
    edited.save(output / "report.docx")
    assert Document(output / "source.docx").paragraphs[1].text == "Jakarta: draft"
    assert Document(output / "report.docx").paragraphs[1].text == "Jakarta: verified"

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
    box.text = "Atlas daily report"
    deck.save(output / "report.pptx")
    reopened = Presentation(output / "report.pptx")
    assert reopened.slides[0].shapes[0].text == "Atlas daily report"

    with (output / "source.csv").open("w", newline="", encoding="utf-8") as handle:
        csv.writer(handle).writerows([['id', 'amount'], ['00123', '9007199254740993']])
    frame = pd.read_csv(output / "source.csv", dtype=str, keep_default_na=False)
    assert frame.iloc[0].tolist() == ['00123', '9007199254740993']
    frame.to_excel(output / "data.xlsx", index=False)
    workbook = load_workbook(output / "data.xlsx")
    assert workbook.active['A2'].value == '00123'
    assert workbook.active['B2'].value == '9007199254740993'
    workbook.active['C1'] = 'reviewed'
    workbook.save(output / "data.xlsx")
    assert load_workbook(output / "data.xlsx").active['C1'].value == 'reviewed'

    pdf = canvas.Canvas(str(output / "report.pdf"))
    for title in ["Atlas first page", "Atlas second page"]:
        pdf.drawString(60, 780, title)
        pdf.showPage()
    pdf.save()
    reader = PdfReader(output / "report.pdf")
    assert len(reader.pages) == 2
    merged = PdfWriter()
    merged.append(reader)
    merged.append(reader, pages=[1])
    merged.write(output / "merged.pdf")
    assert len(PdfReader(output / "merged.pdf").pages) == 3
    split = PdfWriter()
    split.add_page(PdfReader(output / "merged.pdf").pages[2])
    split.write(output / "split.pdf")
    assert "Atlas second page" in PdfReader(output / "split.pdf").pages[0].extract_text()

    print(json.dumps({
        "status": "passed",
        "formats": ["docx", "pptx", "xlsx", "csv", "pdf", "jpg", "png"],
        "versions": {name: importlib.metadata.version(name) for name in
                     ["Pillow", "pandas", "openpyxl", "python-docx", "python-pptx", "pypdf", "reportlab", "fonttools", "uharfbuzz", "python-bidi"]},
    }))


if __name__ == "__main__":
    verify()

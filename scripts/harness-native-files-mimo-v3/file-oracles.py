"""Read native artifacts independently; never expose oracle data to the agent.

Code execution is deliberately separate: supply hidden-test results from an
isolated executor, not from this read-only artifact inspector.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import re
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader


def safe_file(workspace: Path, relative: str) -> Path:
    relative_path = Path(relative)
    if relative_path.is_absolute() or ".." in relative_path.parts:
        raise ValueError("Artifact path must be relative without traversal")
    workspace = workspace.resolve(strict=True)
    path = workspace / relative_path
    if not path.resolve().is_relative_to(workspace):
        raise ValueError("Artifact escapes the trial workspace")
    parent = path
    while parent != workspace:
        if parent.is_symlink():
            raise ValueError("Symlinks are not accepted as artifact evidence")
        parent = parent.parent
    if not path.is_file() or path.stat().st_size > 20_000_000:
        raise ValueError("Artifact is missing or exceeds the 20 MB inspection limit")
    return path


def bounded_archive(path: Path) -> None:
    """Bound decompression before any third-party Office parser opens the ZIP."""
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > 10_000 or sum(item.file_size for item in entries) > 60_000_000:
            raise ValueError("Office archive exceeds inspection bounds")
        if len({item.filename for item in entries}) != len(entries):
            raise ValueError("Office archive contains ambiguous duplicate entries")


def color_value(color) -> tuple:
    kind = color.type
    value = color.rgb if kind == 1 else color.theme_color if kind == 2 else None
    return str(kind), str(value)


def bounded_workbook(path: Path):
    bounded_archive(path)
    workbook = load_workbook(path, data_only=False)
    if len(workbook.worksheets) > 100 or any(sheet.max_row > 20_000 or sheet.max_column > 512 or sheet.max_row * sheet.max_column > 1_000_000 for sheet in workbook):
        raise ValueError("Workbook dimensions exceed inspection bounds")
    return workbook


def verified_original(task: dict, originals: Path) -> Path:
    path = safe_file(originals, task["source"])
    if hashlib.sha256(path.read_bytes()).hexdigest() != task["sources"][task["source"]]:
        raise ValueError("Evaluator original does not match the preregistered source hash")
    return path


def workbook_snapshot(path: Path) -> dict:
    workbook = bounded_workbook(path)
    result = {}
    for sheet in workbook:
        cells = {}
        for row in sheet:
            for cell in row:
                if cell.value is not None or cell.has_style:
                    cells[cell.coordinate] = {
                        "value": cell.value,
                        "type": cell.data_type,
                        "numberFormat": cell.number_format,
                        "font": str(cell.font),
                        "fill": str(cell.fill),
                        "alignment": str(cell.alignment),
                        "border": str(cell.border),
                        "protection": str(cell.protection),
                    }
        result[sheet.title] = {
            "cells": cells,
            "state": sheet.sheet_state,
            "freeze": str(sheet.freeze_panes),
            "merged": sorted(str(item) for item in sheet.merged_cells),
            "columns": {key: (value.width, value.hidden) for key, value in sheet.column_dimensions.items()},
            "rows": {key: (value.height, value.hidden) for key, value in sheet.row_dimensions.items()},
        }
    return {"order": workbook.sheetnames, "sheets": result}


def rich_paragraph(paragraph) -> dict:
    spans = []
    for run in paragraph.runs:
        style = {
            "bold": run.bold,
            "italic": run.italic,
            "underline": run.underline,
            "font": run.font.name,
            "size": run.font.size,
            "color": color_value(run.font.color),
        }
        if not run.text:
            continue
        if spans and spans[-1]["style"] == style:
            spans[-1]["text"] += run.text
        else:
            spans.append({"text": run.text, "style": style})
    return {"text": paragraph.text, "style": paragraph.style.name, "spans": spans}


def document_snapshot(path: Path) -> dict:
    bounded_archive(path)
    document = Document(path)
    paragraphs = [*document.paragraphs, *[p for table in document.tables for row in table.rows for cell in row.cells for p in cell.paragraphs]]
    styles = {}
    for paragraph in paragraphs:
        for initial in [paragraph.style, *[run.style for run in paragraph.runs]]:
            style = initial
            while style is not None and style.style_id not in styles:
                styles[style.style_id] = ET.canonicalize(style.element.xml, strip_text=True, rewrite_prefixes=True)
                style = style.base_style
    return {
        "styles": styles,
        "themes": archive_theme_parts(path, "word/theme/"),
        "paragraphs": [rich_paragraph(p) for p in document.paragraphs],
        "tables": [[[{"text": cell.text, "paragraphs": [rich_paragraph(p) for p in cell.paragraphs]} for cell in row.cells] for row in table.rows] for table in document.tables],
    }


def archive_theme_parts(path: Path, prefix: str) -> dict:
    with zipfile.ZipFile(path) as archive:
        return {name: ET.canonicalize(archive.read(name).decode(), strip_text=True, rewrite_prefixes=True) for name in archive.namelist() if name.startswith(prefix) and name.endswith(".xml")}


def presentation_snapshot(path: Path) -> dict:
    bounded_archive(path)
    presentation = Presentation(path)
    slides = []
    for slide in presentation.slides:
        shapes = []
        for shape in slide.shapes:
            paragraphs = []
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    runs = []
                    for run in paragraph.runs:
                        if not run.text:
                            continue
                        style = {"bold": run.font.bold, "italic": run.font.italic, "underline": run.font.underline, "font": run.font.name, "size": run.font.size,
                                 "color": color_value(run.font.color)}
                        if runs and runs[-1]["style"] == style:
                            runs[-1]["text"] += run.text
                        else:
                            runs.append({"text": run.text, "style": style})
                    paragraphs.append({
                        "text": paragraph.text,
                        "level": paragraph.level,
                        "alignment": paragraph.alignment,
                        "runs": runs,
                    })
            shapes.append({"kind": shape.shape_type, "position": [shape.left, shape.top, shape.width, shape.height], "paragraphs": paragraphs})
        slides.append(shapes)
    return {"slides": slides, "themes": archive_theme_parts(path, "ppt/theme/")}


def replace_text(value: object, old: str, new: str) -> object:
    if isinstance(value, str):
        return value.replace(old, new)
    if isinstance(value, list):
        return [replace_text(item, old, new) for item in value]
    if isinstance(value, dict):
        return {key: replace_text(item, old, new) for key, item in value.items()}
    return value


def contains_fact(text: str, value: object) -> bool:
    if re.fullmatch(r"-?\d+", str(value)):
        # A box/order count must not match a date component, decimal or grouped number.
        return bool(re.search(r"(?<![\w.,+-])" + re.escape(str(value)) + r"(?!\w|[.,-]\d)", text))
    return bool(re.search(r"(?<!\w)" + re.escape(str(value)) + r"(?!\w)", text))


def forbidden_positive_phrases(text: str, pattern: str) -> list[str]:
    """Limited diagnostic, not a general semantic hallucination detector."""
    found = []
    for clause in re.split(r"[.;\n]", text):
        for match in re.finditer(pattern, clause, re.I):
            prefix = clause[:match.start()]
            if re.search(r"\b(no|not|never|unknown|unspecified)\b(?:\W+\w+){0,4}\W*$", prefix, re.I):
                continue
            found.append(match.group())
    return found


def inspect(task: dict, workspace: Path, originals: Path, hidden_code: dict | None = None,
            delivered_path: str | None = None, expected_sandbox_policy_sha256: str | None = None,
            expected_candidate_contract_sha256: str | None = None) -> dict:
    checks = []

    def check(name: str, passed: bool, detail: str = "") -> None:
        checks.append({"id": name, "pass": bool(passed), "detail": detail})

    integrity = False
    for relative, expected_hash in task["sources"].items():
        if relative in task.get("mutableSources", []):
            continue
        try:
            original = safe_file(workspace, relative)
            preserved = hashlib.sha256(original.read_bytes()).hexdigest() == expected_hash
            check(f"source_preserved:{relative}", preserved)
            integrity |= not preserved
        except (OSError, ValueError) as error:
            check(f"source_preserved:{relative}", False, str(error))
            integrity = True
    try:
        if task["kind"] == "code_fix":
            selected = task["output"]
        else:
            if not isinstance(delivered_path, str) or not delivered_path:
                raise ValueError("Delivery unresolved: select one path from final delivery evidence before scoring")
            selected = delivered_path
            relative = Path(selected)
            if not relative.parts or relative.parts[0] != "artifacts" or relative.suffix.lower() != Path(task["output"]).suffix.lower():
                raise ValueError("Selected delivery must be under artifacts/ with the required extension")
        check("delivery_selected", True, selected)
        output = safe_file(workspace, selected)
        check("deliverable_exists", True)
        kind, expected = task["kind"], task["expected"]
        if kind == "csv":
            rows = list(csv.reader(io.StringIO(output.read_text(encoding="utf-8-sig"), newline="")))
            check("csv_content", rows == expected)
        elif kind == "xlsx_cells":
            workbook = bounded_workbook(output)
            check("sheet_names", workbook.sheetnames == list(expected))
            for name, rows in expected.items():
                if name not in workbook:
                    check(f"cells:{name}", False, "Required sheet missing")
                    continue
                sheet = workbook[name]
                actual = [[sheet.cell(r + 1, c + 1).value for c in range(len(row))] for r, row in enumerate(rows)]
                check(f"cells:{name}", actual == rows)
                check(f"schema_extra_cells:{name}", all(cell.value is None or (cell.row <= len(rows) and cell.column <= len(rows[cell.row - 1])) for row in sheet for cell in row))
                check(f"cell_types:{name}", all(isinstance(workbook[name].cell(r + 1, c + 1).value, (int, float)) and not isinstance(workbook[name].cell(r + 1, c + 1).value, bool) for r, row in enumerate(rows) for c, value in enumerate(row) if isinstance(value, int)))
        elif kind == "xlsx_edit":
            source = workbook_snapshot(verified_original(task, originals))
            source["sheets"][expected["sheet"]]["cells"][expected["cell"]]["value"] = expected["value"]
            check("workbook_edit_and_preservation", workbook_snapshot(output) == source)
        elif kind == "docx_edit":
            source = document_snapshot(verified_original(task, originals))
            wanted = replace_text(source, expected["old"], expected["new"])
            check("document_edit_and_preservation", document_snapshot(output) == wanted)
        elif kind == "docx_report":
            bounded_archive(output)
            document = Document(output)
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            headings = [p for p in document.paragraphs if p.style.name == "Title" or p.style.name.startswith("Heading ")]
            check("title_heading", bool(headings) and headings[0].text == "Incident report")
            for name, value in expected.items():
                check(f"fact:{name}", contains_fact(text, value))
            tables = [[[cell.text for cell in row.cells] for row in table.rows] for table in document.tables]
            check("action_table", tables == [[["Action", "Due"], [expected["action"], expected["due"]]]])
            claims = forbidden_positive_phrases(text, r"\b(resolved|resolution completed|loss of|duration of)\b")
            check("forbidden_positive_phrase", not claims, "; ".join(claims))
        elif kind == "pdf":
            reader = PdfReader(output)
            text = " ".join(" ".join((page.extract_text() or "").split()) for page in reader.pages)
            check("page_count", len(reader.pages) == expected["pages"])
            for index, value in enumerate(expected["text"]):
                check(f"fact:{index}", contains_fact(text, value))
            claims = forbidden_positive_phrases(text, r"\b(approved by|signed by|delivered successfully)\b")
            check("forbidden_positive_phrase", not claims, "; ".join(claims))
        elif kind == "pptx_edit":
            source = presentation_snapshot(verified_original(task, originals))
            check("deck_edit_and_preservation", presentation_snapshot(output) == replace_text(source, expected["old"], expected["new"]))
        elif kind == "code_fix":
            extra_source_files = [str(path.relative_to(workspace)) for path in (workspace / "app").rglob("*") if path.is_file() and str(path.relative_to(workspace)) not in task["sources"] and not (path.parent.name == "__pycache__" and path.suffix == ".pyc")]
            check("code_mutation_scope", not extra_source_files, ", ".join(extra_source_files))
            integrity |= bool(extra_source_files)
            receipt = hidden_code or {}
            count = len(expected["cases"]) + len(expected["invalid"])
            valid_policy = isinstance(expected_sandbox_policy_sha256, str) and bool(re.fullmatch(r"[a-f0-9]{64}", expected_sandbox_policy_sha256))
            valid_contract = isinstance(expected_candidate_contract_sha256, str) and bool(re.fullmatch(r"[a-f0-9]{64}", expected_candidate_contract_sha256))
            admission = receipt.get("candidateContract")
            check("isolated_hidden_tests", bool(
                receipt.get("pass") is True
                and receipt.get("sourceSha256") == hashlib.sha256(output.read_bytes()).hexdigest()
                and receipt.get("testSuiteSha256") == task["testSuiteSha256"]
                and valid_policy and receipt.get("sandboxPolicySha256") == expected_sandbox_policy_sha256
                and valid_contract
                and task.get("candidateContractSha256") == expected_candidate_contract_sha256
                and isinstance(admission, dict) and admission.get("pass") is True
                and admission.get("sha256") == expected_candidate_contract_sha256
                and type(receipt.get("exitCode")) is int and receipt["exitCode"] == 0
                and type(receipt.get("casesTotal")) is int and receipt["casesTotal"] == count
                and type(receipt.get("casesPassed")) is int and receipt["casesPassed"] == count
            ), "Requires an orchestrator-owned isolated execution receipt bound to source, suite, independently expected sandbox policy and disclosed candidate contract, successful admission, exit status and all cases.")
        else:
            raise ValueError("Unknown oracle kind")
    except Exception as error:
        check("artifact_inspection", False, f"{type(error).__name__}: {error}")
    return {"pass": all(item["pass"] for item in checks), "integrityFailure": integrity, "checks": checks, "scope": "Native artifact structure and required supplied facts; forbidden positive phrases are a limited diagnostic, not comprehensive supported-fact proof. No claim of visual perfection, OCR quality, or general code correctness."}


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("manifest", type=Path)
    parser.add_argument("workspace", type=Path)
    parser.add_argument("originals", type=Path)
    parser.add_argument("--hidden-code", type=Path)
    parser.add_argument("--delivered-path")
    parser.add_argument("--sandbox-policy-sha256")
    parser.add_argument("--candidate-contract-sha256")
    args = parser.parse_args()
    hidden = json.loads(args.hidden_code.read_text()) if args.hidden_code else None
    print(json.dumps(inspect(json.loads(args.manifest.read_text()), args.workspace, args.originals, hidden, args.delivered_path, args.sandbox_policy_sha256, args.candidate_contract_sha256), indent=2))

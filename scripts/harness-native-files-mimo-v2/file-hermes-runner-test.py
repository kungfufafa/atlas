#!/usr/bin/env python3
"""Offline native integration gates. Run with the pinned Hermes venv on macOS.

All HTTP is fake loopback. Generated commands contain synthetic data only;
no provider key, user configuration or candidate implementation is executed.
"""

import contextlib
import importlib.util
import json
import os
from pathlib import Path
import shlex
import subprocess
import sys
import tempfile
import threading
import time
import unittest
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

RUNNER = Path(__file__).with_name("file_hermes_runner.py")
SPEC = importlib.util.spec_from_file_location("hermes_file_runner", RUNNER)
adapter = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(adapter)


class FakeProvider:
    def __init__(self, script, *, freeze_after_tool=False):
        self.requests = []
        self.script = script
        self.freeze = freeze_after_tool
        self.invoked = False
        self.release = threading.Event()
        owner = self

        class Handler(BaseHTTPRequestHandler):
            def log_message(self, *args):
                pass

            def reply(self, body):
                encoded = json.dumps(body).encode()
                with contextlib.suppress(BrokenPipeError, ConnectionResetError):
                    self.send_response(200)
                    self.send_header("Content-Type", "application/json")
                    self.send_header("Content-Length", str(len(encoded)))
                    self.end_headers()
                    self.wfile.write(encoded)

            def do_GET(self):
                self.reply({"data": [{"id": "offline-native-file"}]})

            def do_POST(self):
                request = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
                owner.requests.append(request)
                names = {tool["function"]["name"] for tool in request.get("tools", [])}
                tool = None
                if names and not owner.invoked:
                    if "terminal" in names:
                        owner.invoked = True
                        tool = {"name": "terminal", "arguments": json.dumps(owner.script)}
                    else:
                        tool = {"name": "tool_search", "arguments": json.dumps({"query": "terminal"})}
                elif names and owner.freeze:
                    owner.release.wait(20)
                message = {"role": "assistant", "content": "[Workbook](artifacts/result.xlsx)" if names else '{"title":"Offline file proof"}'}
                if tool:
                    message = {"role": "assistant", "content": "", "tool_calls": [{"id": f"native-{len(owner.requests)}", "type": "function", "function": tool}]}
                self.reply({"id": "offline", "model": "offline-native-file", "choices": [{"index": 0, "finish_reason": "tool_calls" if tool else "stop", "message": message}], "usage": {"prompt_tokens": 20, "completion_tokens": 10, "total_tokens": 30}})

        self.server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
        self.thread = threading.Thread(target=self.server.serve_forever, daemon=True)
        self.thread.start()

    @property
    def url(self):
        return f"http://127.0.0.1:{self.server.server_port}"

    def close(self):
        self.release.set()
        self.server.shutdown()
        self.server.server_close()
        self.thread.join(timeout=2)


class RunnerTests(unittest.TestCase):
    def setUp(self):
        self.root = Path(tempfile.mkdtemp(prefix="hermes-file-gate-", dir=adapter.ROOT))
        self.source = self.root / "source"
        (self.source / "input").mkdir(parents=True)
        (self.source / "input" / "note.txt").write_text("synthetic fixture")

    def request(self, origin="http://127.0.0.1:43199", timeout=30000):
        return {"runId": "offline-file-gate", "proxyBaseUrl": origin, "model": "offline-native-file", "budget": {"maxProviderRequests": 24, "maxGeneratedTokens": 12000, "maxOutputTokens": 4096, "timeoutMs": timeout}, "taskTurns": ["Use the prepared Python interpreter to create the requested file, then return its link."], "sourceDirectory": str(self.source), "stateRoot": str(self.root / "state"), "pythonPath": str(adapter.PREPARED_PYTHON)}

    def invoke(self, provider, timeout=30000):
        request = self.request(provider.url, timeout)
        result = subprocess.run([sys.executable, "-B", str(RUNNER)], input=json.dumps(request), text=True, capture_output=True, timeout=timeout / 1000 + 15, env={"PATH": "/opt/homebrew/bin:/usr/bin:/bin", "TMPDIR": "/private/tmp", "FILE_ADAPTER_TEST_SENTINEL": "synthetic-secret"})
        (self.root / "stdout.json").write_text(result.stdout)
        (self.root / "stderr.log").write_text(result.stderr)
        (self.root / "requests.json").write_text(json.dumps(provider.requests))
        self.assertEqual(result.returncode, 0, result.stderr[-2000:])
        return json.loads(result.stdout)

    def test_rejects_arbitrary_runtime_and_runtime_alias_cannot_widen_policy(self):
        request = self.request()
        request["pythonPath"] = "/bin/true"
        with self.assertRaises(ValueError):
            adapter.check_request(request)
        alias = self.root / "python"
        alias.symlink_to(adapter.PREPARED_PYTHON)
        request["pythonPath"] = str(alias)
        adapter.check_request(request)
        profile = adapter.policy(self.root / "workspace", self.root / "native", alias, self.root / "worker.py", 43199)
        self.assertNotIn(f'(subpath "{self.root.parent}")', profile)
        self.assertNotIn('(subpath "/")', profile)
        self.assertIn(str(adapter.PREPARED_PYTHON.parent.parent), profile)

    def test_request_blocks_expectations_remote_endpoint_and_nonpositive_budget(self):
        for mutation in ({"expected": {}}, {"proxyBaseUrl": "https://example.com:443"}, {"taskTurns": []}):
            with self.assertRaises(ValueError):
                adapter.check_request({**self.request(), **mutation})
        for value in (0, -1, True):
            request = self.request()
            request["budget"]["timeoutMs"] = value
            with self.assertRaises(ValueError):
                adapter.check_request(request)

    def test_source_copy_refuses_symlink_and_evaluator_file(self):
        (self.source / "expected.json").write_text("{}")
        with self.assertRaises(ValueError):
            adapter.copy_sources(self.source, self.root / "workspace")
        (self.source / "expected.json").unlink()
        (self.source / "input" / "alias").symlink_to(self.root / "outside")
        with self.assertRaises(ValueError):
            adapter.copy_sources(self.source, self.root / "workspace2")

    def test_stock_terminal_writes_xlsx_with_host_and_network_canaries_blocked(self):
        outside = self.root / "outside.txt"
        outside.write_text("synthetic outside canary")
        denied = FakeProvider({})
        child = f"from pathlib import Path\ntry:\n Path({str(outside)!r}).read_text()\n print('ALLOWED')\nexcept OSError:\n print('blocked')"
        program = f"""import json, os, socket, subprocess, sys
from pathlib import Path
from openpyxl import Workbook
proof = {{}}
outside = Path({str(outside)!r})
Path('escape-link').symlink_to(outside)
for name, action in [('outside_read', lambda: outside.read_text()), ('outside_write', lambda: outside.write_text('changed')), ('symlink_read', lambda: Path('escape-link').read_text()), ('network', lambda: socket.create_connection(('127.0.0.1', {denied.server.server_port}), timeout=1))]:
 try:
  action()
  proof[name] = 'ALLOWED'
 except OSError:
  proof[name] = 'blocked'
proof['descendant_read'] = subprocess.run([sys.executable, '-I', '-c', {child!r}], capture_output=True, text=True, timeout=3).stdout.strip()
proof['environment_scrubbed'] = 'FILE_ADAPTER_TEST_SENTINEL' not in os.environ
Path('artifacts').mkdir()
w = Workbook()
w.active.append(['id', 'cents'])
w.active.append(['0007', 1234])
w.save('artifacts/result.xlsx')
Path('artifacts/proof.json').write_text(json.dumps(proof))
print(json.dumps(proof))
"""
        provider = FakeProvider({"command": shlex.join([str(adapter.PREPARED_PYTHON), "-I", "-c", program]), "timeout": 20})
        try:
            result = self.invoke(provider)
            self.assertEqual(result["status"], "completed")
            workspace = Path(result["workspaceRoot"])
            proof = json.loads((workspace / "artifacts/proof.json").read_text())
            self.assertEqual(proof, {"outside_read": "blocked", "outside_write": "blocked", "symlink_read": "blocked", "network": "blocked", "descendant_read": "blocked", "environment_scrubbed": True})
            self.assertEqual(outside.read_text(), "synthetic outside canary")
            self.assertEqual(denied.requests, [])
            rows = subprocess.check_output([str(adapter.PREPARED_PYTHON), "-I", "-c", "import json,sys;from openpyxl import load_workbook;print(json.dumps(list(load_workbook(sys.argv[1]).active.values)))", str(workspace / "artifacts/result.xlsx")], text=True)
            self.assertEqual(json.loads(rows), [["id", "cents"], ["0007", 1234]])
            receipts = [event for event in result["nativeEvents"] if event["name"] == "terminal"]
            self.assertEqual(len(receipts), 1)
            self.assertEqual(json.loads(receipts[0]["result"])["exit_code"], 0)
            self.assertTrue(result["evidence"]["artifactInspectionAllowed"])
            self.assertTrue(result["evidence"]["nativeCleanup"]["completed"])
            self.assertEqual(result["evidence"]["supervisorCleanup"]["remainingPids"], [])
            self.assertFalse(result["evidence"]["caBundle"]["verificationDisabled"])
        finally:
            provider.close()
            denied.close()

    def background_case(self, freeze):
        import psutil
        program = "import os,time;from pathlib import Path;Path('background.pid').write_text(str(os.getpid()));time.sleep(120)"
        provider = FakeProvider({"command": shlex.join([str(adapter.PREPARED_PYTHON), "-I", "-c", program]), "background": True}, freeze_after_tool=freeze)
        started = time.monotonic()
        try:
            result = self.invoke(provider, timeout=8000 if freeze else 30000)
            self.assertEqual(result["status"], "budget_exceeded" if freeze else "completed")
            self.assertTrue(any(event["name"] == "terminal" for event in result["nativeEvents"]))
            self.assertTrue(result["evidence"]["nativeCleanup"]["completed"])
            self.assertEqual(result["evidence"]["supervisorCleanup"]["remainingPids"], [])
            self.assertFalse(result["evidence"]["workerHardKilled"])
            self.assertEqual(result["evidence"]["nativeCleanup"]["isolatedRegistryKilled"], 1)
            receipts = [event for event in result["nativeEvents"] if event["name"] == "terminal"]
            pids = [json.loads(receipts[0]["result"])["pid"]]
            pid_path = Path(result["workspaceRoot"]) / "background.pid"
            # Fast successful shutdown can kill the stock shell before Python
            # writes readiness. The native spawn receipt still identifies it.
            if pid_path.exists():
                pids.append(int(pid_path.read_text()))
            for pid in pids:
                with contextlib.suppress(psutil.NoSuchProcess):
                    self.assertEqual(psutil.Process(pid).status(), psutil.STATUS_ZOMBIE)
            if freeze:
                self.assertLess(time.monotonic() - started, 17)
            else:
                self.assertTrue(result["evidence"]["artifactInspectionAllowed"])
        finally:
            provider.close()

    def test_worker_shutdown_stops_stock_background_terminal_before_inspection(self):
        self.background_case(False)

    def test_shared_deadline_preserves_receipts_and_stops_stock_background_terminal(self):
        self.background_case(True)


if __name__ == "__main__":
    unittest.main(verbosity=2)

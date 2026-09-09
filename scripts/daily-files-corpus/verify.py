"""Independent oracle. Appends every attempt; refusals never become passes."""
import argparse
import csv
import hashlib
import json
import re
import zipfile
from pathlib import Path
from datetime import datetime
from docx import Document
from lxml import etree
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader


def sha(path): return hashlib.sha256(path.read_bytes()).hexdigest()


def xml_structure(node):
    # Text changes are allowed; styles, relationships, geometry and other
    # attributes/elements inside an edited XML part must remain unchanged.
    is_text = etree.QName(node).localname == "t"
    attributes = sorted((key,value) for key,value in node.attrib.items() if key != "{http://www.w3.org/XML/1998/namespace}space")
    return (node.tag,attributes,"" if is_text else (node.text or "").strip(),[xml_structure(child) for child in node if isinstance(child.tag,str)])



def verify(root, execution_path, revision):
    revision_preservation = revision not in {"v1", "v2", "v3", "v4", "v5"}
    manifest = json.loads((root/"manifest.json").read_text()); execution = json.loads(execution_path.read_text())
    cases = {entry["id"]:entry for entry in manifest["cases"]}; results = []
    for record in execution["records"]:
        case_id = record["id"]; case = cases[case_id]; folder=Path(record["directory"]); checks=[]
        def check(name, condition, detail=None): checks.append({"check":name,"ok":bool(condition),"detail":detail})
        for source in case["sources"]: check("source preserved: "+source["path"],sha(folder/source["path"]) == source["sha256"])
        def output(tool,operation=None):
            candidates=[step["output"] for step in record["steps"] if step["tool"]==tool and "output" in step and (operation is None or step["input"].get("operation",step["input"].get("action"))==operation)]
            if not candidates: raise ValueError("Missing output for "+tool+" "+str(operation))
            return candidates[-1]
        def artifact(tool,operation=None): return folder/output(tool,operation)["path"]
        status=None; reason=record.get("error"); detail={}
        if reason:
            refused = re.search(r"unsupported|cannot preserve|cannot be converted losslessly|exceeds the .*byte limit|cannot render|requires LibreOffice|needs OCR|unmerged cell",reason,re.I)
            status="unsupported" if refused else "failure"
        try:
            if not reason and case_id[0] in {"D","P"}:
                extension="docx" if case_id.startswith("D") else "pptx"
                destination=artifact("office_document","edit"); allowed=set(case["expected"]["allowedParts"])
                with zipfile.ZipFile(folder/("input."+extension)) as before, zipfile.ZipFile(destination) as after:
                    check("same package parts",set(before.namelist())==set(after.namelist()))
                    changed=[name for name in before.namelist() if before.read(name)!=after.read(name)]
                    check("only declared XML parts changed",set(changed).issubset(allowed),changed)
                    for name in changed:
                        if name in allowed and not (case_id=="D03" and revision_preservation):
                            check("edited part retains non-text XML structure: "+name,xml_structure(etree.fromstring(before.read(name)))==xml_structure(etree.fromstring(after.read(name))))
                if case_id=="D01":
                    document=Document(destination); check("multilingual replacement",document.paragraphs[0].text==case["expected"]["text"],document.paragraphs[0].text)
                    check("run styles retained",[(run.bold,run.italic,run.underline) for run in document.paragraphs[0].runs]==[(True,None,None),(None,True,None),(None,None,True)])
                    check("image retained",len(document.inline_shapes)==1)
                elif case_id=="D02":
                    table=Document(destination).tables[0]; check("inner cell changed",table.cell(1,1).tables[0].cell(1,1).text=="Inner approved")
                    check("outer cells retained",table.cell(1,0).text=="Outer keeper" and table.cell(0,1).text=="Unchanged")
                elif case_id=="D03":
                    with zipfile.ZipFile(destination) as archive: xml=etree.fromstring(archive.read("word/document.xml"))
                    check("tracked insertion changed coherently","Inserted approved" in "".join(xml.itertext()))
                    if revision_preservation:
                        from copy import deepcopy
                        ns={"w":"http://schemas.openxmlformats.org/wordprocessingml/2006/main"}; w="{"+ns["w"]+"}"
                        with zipfile.ZipFile(folder/"input.docx") as archive: original=etree.fromstring(archive.read("word/document.xml"))
                        originals=original.xpath("//w:ins|//w:del",namespaces=ns)
                        revisions=xml.xpath("//w:ins|//w:del",namespaces=ns)
                        fresh=[node for node in revisions if node.get(w+"author")=="Atlas"]
                        check("original reviewer identities retained",all(any(node.tag==old.tag and node.attrib==old.attrib for node in revisions) for old in originals))
                        ids=[node.get(w+"id") for node in revisions]
                        check("new revision ids unique",len(ids)==len(set(ids)) and all(value is not None for value in ids))
                        check("new insert and deletion attributed and dated",{etree.QName(node).localname for node in fresh}=={"ins","del"} and all(node.get(w+"date") for node in fresh))
                        def resolve_atlas(accept):
                            result=deepcopy(xml)
                            for node in list(result.xpath("//w:ins[@w:author='Atlas']|//w:del[@w:author='Atlas']",namespaces=ns)):
                                parent=node.getparent()
                                if parent is None: continue
                                if (etree.QName(node).localname=="ins") != accept: parent.remove(node); continue
                                for text in node.xpath(".//w:delText",namespaces=ns): text.tag=w+"t"
                                index=parent.index(node)
                                for child in list(node): parent.insert(index,child); index+=1
                                parent.remove(node)
                            return result
                        def exact_tree(node):
                            return (node.tag,sorted((key,value) for key,value in node.attrib.items() if key!="{http://www.w3.org/XML/1998/namespace}space"),node.text or "",[exact_tree(child) for child in node])
                        check("rejecting only Atlas revisions restores original document",exact_tree(resolve_atlas(False))==exact_tree(original))
                        accepted=resolve_atlas(True)
                        check("accepting Atlas revisions yields requested visible text","".join(accepted.xpath("//w:t/text()",namespaces=ns))=="Inserted approved")
                elif case_id=="D04":
                    document=Document(destination); check("second header changed",document.sections[1].header.paragraphs[0].text=="Second section approved")
                    check("first header retained",document.sections[0].header.paragraphs[0].text=="First section keep")
                else:
                    presentation=Presentation(destination); slide=presentation.slides[0]
                    def texts(shapes):
                        values=[]
                        for shape in shapes:
                            if shape.has_text_frame: values.append(shape.text)
                            if hasattr(shape,"shapes"): values.extend(texts(shape.shapes))
                        return values
                    values=texts(slide.shapes)
                    if case_id=="P01":
                        check("caption changed","Quarter approved" in values); chart=next(shape.chart for shape in slide.shapes if shape.has_chart)
                        check("chart series retained",list(chart.series[0].values)==[12,8,15])
                    elif case_id=="P02":
                        check("group member changed","Grouped approved" in values); check("other member retained","Grouped keeper" in values)
                        original=Presentation(folder/"input.pptx").slides[0]
                        check("group and shape bounds retained",[(s.left,s.top,s.width,s.height) for s in slide.shapes]==[(s.left,s.top,s.width,s.height) for s in original.shapes])
                    elif case_id=="P03":
                        table=next(shape.table for shape in slide.shapes if shape.has_table); check("merged content changed",table.cell(0,0).text=="Merged approved"); check("merge retained",table.cell(0,0).is_merge_origin)
                    elif case_id=="P04":
                        check("multilingual text changed",case["expected"]["text"] in values,values)
                        paragraph=next(shape.text_frame.paragraphs[0] for shape in slide.shapes if shape.has_text_frame)
                        check("RTL metadata retained",paragraph._p.get_or_add_pPr().get("rtl")=="1")
            elif not reason and case_id.startswith("X"):
                destination=artifact("spreadsheet","recalculate" if case_id in {"X02","X04"} else "write_range")
                workbook=load_workbook(destination,data_only=False); values=load_workbook(destination,data_only=True)
                if case_id=="X01":
                    sheet=workbook["Records"]; original=load_workbook(folder/"input.xlsx")
                    check("input changed",sheet["C2"].value==11); check("ID remains text",sheet["A2"].value=="000042" and sheet["A2"].data_type=="s")
                    check("date and formats retained",sheet["B2"].value==datetime(2026,9,6) and sheet["B2"].number_format=="yyyy-mm-dd" and sheet["C2"].number_format=="0.00")
                    check("hidden sheet retained",workbook["Lookup"].sheet_state=="hidden" and workbook["Lookup"]["B1"].value==99)
                    check("comment retained",sheet["A2"].comment is not None and sheet["A2"].comment.text==original["Records"]["A2"].comment.text and sheet["A2"].comment.author==original["Records"]["A2"].comment.author)
                    check("validation retained",len(sheet.data_validations.dataValidation)==1 and str(sheet.data_validations.dataValidation[0].sqref)=="D2:D20")
                    check("hyperlink and freeze panes retained",sheet["D2"].hyperlink is not None and sheet["D2"].hyperlink.target=="https://example.invalid/state" and sheet.freeze_panes=="B2")
                elif case_id=="X02":
                    actual=[values["Summary"][cell].value for cell in ["B2","B3","B4"]]
                    check("cross-sheet calculated values",actual==[19,"Review",datetime(2026,9,6)],str(actual)); check("formulas retained",all(workbook["Summary"][cell].data_type=="f" for cell in ["B2","B3","B4"]))
                elif case_id=="X03":
                    check("chart and source edit",workbook["Chart data"]["B2"].value==12 and len(workbook["Chart data"]._charts)==1)
                    if revision_preservation:
                        with zipfile.ZipFile(folder/"input.xlsx") as before, zipfile.ZipFile(destination) as after:
                            check("chart workbook package parts retained",set(before.namelist())==set(after.namelist()))
                            chart_parts=[name for name in before.namelist() if name.startswith("xl/charts/") and name.endswith(".xml")]
                            for name in chart_parts:
                                source_chart=etree.fromstring(before.read(name)); output_chart=etree.fromstring(after.read(name))
                                for root_chart in [source_chart,output_chart]:
                                    for cache in root_chart.xpath("//*[local-name()='numCache' or local-name()='strCache' or local-name()='multiLvlStrCache']"): cache.getparent().remove(cache)
                                check("chart series and formatting retained: "+name,xml_structure(source_chart)==xml_structure(output_chart))
                            for name in before.namelist():
                                if name not in chart_parts and name not in {"xl/workbook.xml","xl/worksheets/sheet1.xml"}:
                                    check("unrelated chart-workbook part retained: "+name,before.read(name)==after.read(name))
                        original=load_workbook(folder/"input.xlsx")["Chart data"]
                        check("other chart-source cells retained",all(cell.value==workbook["Chart data"][cell.coordinate].value for row in original for cell in row if cell.coordinate!="B2"))
                elif case_id=="X04":
                    sheet=workbook["Table data"]; check("table retained","TasksTable" in sheet.tables)
                    table=sheet.tables.get("TasksTable"); check("totals definition retained",table is not None and table.totalsRowCount==1)
                    original=load_workbook(folder/"input.xlsx")["Table data"]
                    source_table=original.tables["TasksTable"]
                    check("table range and style retained",table is not None and table.ref==source_table.ref and (table.tableStyleInfo.name,[bool(getattr(table.tableStyleInfo,key)) for key in ["showFirstColumn","showLastColumn","showRowStripes","showColumnStripes"]])==(source_table.tableStyleInfo.name,[bool(getattr(source_table.tableStyleInfo,key)) for key in ["showFirstColumn","showLastColumn","showRowStripes","showColumnStripes"]]))
                    check("unchanged identifiers and input retained",sheet["A2"].value=="0001" and sheet["A3"].value=="0002" and sheet["B3"].value==5)
                    def cf_rules(value): return [(str(key.sqref),[etree.tostring(rule.to_tree()) for rule in rules]) for key,rules in value.conditional_formatting._cf_rules.items()]
                    check("conditional format retained",cf_rules(sheet)==cf_rules(original)); check("structured formula total",values["Table data"]["B4"].value==17,values["Table data"]["B4"].value)
            elif not reason and case_id in {"C01","C03"}:
                exported=artifact("spreadsheet","export_csv")
                with exported.open(newline="",encoding="utf-8-sig") as handle: rows=list(csv.reader(handle,delimiter=case["expected"]["delimiter"]))
                check("parsed records preserved",rows==case["expected"]["rows"],rows)
                imported=load_workbook(artifact("spreadsheet","import_csv")); check("IDs and expressions remain strings",imported.active["A2"].data_type=="s" and imported.active["C2" if case_id=="C03" else "D2"].data_type=="s")
            elif not reason and case_id=="C02":
                workbook=load_workbook(artifact("spreadsheet","import_csv")); check("locale decimal conversion",workbook.active["B2"].value==12.5 and workbook.active["B3"].value==-0.75); check("identifiers remain text",workbook.active["A2"].value=="00012" and workbook.active["A2"].data_type=="s")
            elif not reason and case_id=="F01":
                extracted=output("pdf_document","extract"); pages=extracted["pages"]
                check("selected page order",[page["page"] for page in pages]==[3,2]); check("selected Unicode text",all(text in page["text"] for text,page in zip(case["expected"]["texts"],pages)),[page["text"] for page in pages])
            elif not reason and case_id=="F02":
                extracted=output("pdf_document","extract"); inspected=output("pdf_document","inspect")
                actual={field["name"]:field["value"] for field in inspected.get("form",{}).get("fields",[])}
                missing=[name for name,value in case["expected"]["fields"].items() if actual.get(name)!=value]
                if missing: status="unsupported"; reason="Filled form field values are outside the returned coverage."; detail["missingValues"]=missing
                else: check("filled fields read as named values",actual==case["expected"]["fields"],actual)
            elif not reason and case_id=="F03":
                extracted=output("pdf_document","extract"); text="\n".join(page["text"] for page in extracted["pages"])
                if case["expected"]["text"] not in text: status="unsupported"; reason="Image-only PDF text was not read; no OCR operation completed."; detail["extract"]=extracted
                else: check("scanned text read",True)
            elif not reason and case_id=="F04":
                reader=PdfReader(artifact("pdf_document","merge")); check("merged page count",len(reader.pages)==3)
                check("outline retained",len(reader.outline)>=2,len(reader.outline)); attachments=reader.attachments
                check("attachment retained",case["expected"]["attachment"] in attachments,list(attachments))
                if case["expected"]["attachment"] in attachments: check("attachment bytes retained",attachments[case["expected"]["attachment"]][0].decode()==case["expected"]["attachmentText"])
                destinations=[]
                for annotation in reader.pages[0].get("/Annots",[]):
                    value=annotation.get_object()
                    if "/Dest" in value and isinstance(value["/Dest"],list): destinations.append(reader.get_page_number(value["/Dest"][0].get_object()))
                check("internal destination remains in merged page tree",bool(destinations) and all(isinstance(number,int) and number>=0 for number in destinations),destinations)
            elif not reason and case_id=="F05":
                text="\n".join(page.extract_text() for page in PdfReader(artifact("pdf_document","create")).pages)
                check("multilingual PDF extracted text",case["expected"]["text"] in text,text); detail["visualReview"]="pending"
            elif case_id=="W01":
                if reason: status="failure"
                if execution["attempt"]=="initial-20260906": detail["driverDefect"]="CSV import omitted explicit numeric column type; corrected in subsequent attempt, initial result retained."
                workbook=load_workbook(artifact("spreadsheet","recalculate"),data_only=True); total=workbook["records"]["B5"].value
                check("independent total",total==18,total)
                office_outputs=[step["output"] for step in record["steps"] if step["tool"]=="office_document" and step["input"].get("operation")=="edit" and "output" in step]
                report=Document(folder/office_outputs[0]["path"]); check("DOCX total","Total units: 18" in [p.text for p in report.paragraphs])
                presentation=Presentation(folder/office_outputs[1]["path"]); check("PPTX total",any(shape.has_text_frame and shape.text=="Total units: 18" for shape in presentation.slides[0].shapes))
                packet=PdfReader(artifact("pdf_document","merge")); text="\n".join(page.extract_text() for page in packet.pages)
                check("PDF packet total and appendix","Total units: 18" in text and "Independent appendix retained" in text,text)
        except Exception as error:
            status="failure"; reason="Independent oracle could not validate result: "+str(error)
        failed=[check for check in checks if not check["ok"]]
        if failed: status="failure"; reason=reason or "Required independent preservation/content assertions failed."
        if status is None: status="supported"
        result={"id":case_id,"attempt":execution["attempt"],"oracleRevision":revision,"status":status,"reason":reason,"checks":checks,"detail":detail,"execution":str(execution_path),"durationMs":record["durationMs"]}
        results.append(result)
    ledger=root/"attempts.jsonl"
    existing = [json.loads(line) for line in ledger.read_text().splitlines()] if ledger.exists() else []
    if any(item["attempt"]==execution["attempt"] and item.get("oracleRevision","v1")==revision for item in existing):
        raise ValueError("This attempt/oracle revision already exists; choose a new revision to preserve the evidence ledger.")
    with ledger.open("a") as handle:
        for result in results: handle.write(json.dumps(result,ensure_ascii=False)+"\n")
    report=execution_path.parent/("verified.oracle-"+revision+".json"); report.write_text(json.dumps(results,ensure_ascii=False,indent=2))
    counts={status:sum(item["status"]==status for item in results) for status in ["supported","unsupported","failure"]}
    print(json.dumps({"attempt":execution["attempt"],"counts":counts,"report":str(report)},ensure_ascii=False))
    for result in results:
        print(json.dumps({"id":result["id"],"status":result["status"],"reason":result["reason"],"failedChecks":[item for item in result["checks"] if not item["ok"]]},ensure_ascii=False))


if __name__=="__main__":
    parser=argparse.ArgumentParser(); parser.add_argument("root",type=Path); parser.add_argument("execution",type=Path); parser.add_argument("--revision",default="v6"); args=parser.parse_args(); verify(args.root,args.execution,args.revision)

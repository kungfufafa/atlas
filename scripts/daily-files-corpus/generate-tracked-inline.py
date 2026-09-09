"""Additional Word revision fixture, separate from the unchanged D03 source.

python-docx 1.2.0 with an inline OOXML insertion around the text run. The
original D03 wraps a whole paragraph in an insertion; LibreOffice does not
render that source paragraph, so it is unsuitable as visual evidence.
"""
import sys
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def generate(destination):
    document = Document()
    paragraph = document.add_paragraph()
    insertion = OxmlElement("w:ins")
    insertion.set(qn("w:id"), "4")
    insertion.set(qn("w:author"), "Synthetic reviewer")
    insertion.set(qn("w:date"), "2026-09-06T00:00:00Z")
    run = paragraph.add_run("Inserted draft")
    run.bold = True
    paragraph._p.remove(run._r)
    insertion.append(run._r)
    paragraph._p.append(insertion)
    document.save(destination)


if __name__ == "__main__":
    generate(Path(sys.argv[1]))
